import pytest
import openpyxl
from playwright.sync_api import Page, Playwright, expect

# CREATE EMPTY LIST
login_data=[]

# LOAD WORKBOOK
workbook=openpyxl.load_workbook('D:/PythonProject/testData/data.xlsx')
# LOAD SHEET
sheet=workbook.active

# Read data from excel sheet

for row in sheet.iter_rows(min_row=2, values_only=True): # .iter_rows pre defined function
    username, password, expected_success = row
    login_data.append((str(username or ''), str(password or ''), str(expected_success or ''))) #convert to string and add '' to read blank data from the sheet
workbook.close()

@pytest.mark.parametrize('username,password,expected_success',login_data)
def test_login_json(username, password, expected_success,page:Page):
    page.goto('https://demowebshop.tricentis.com/')
    page.locator('.ico-login').click()
    page.locator('#Email').fill(username)
    page.locator('#Password').fill(password)
    page.locator('.button-1.login-button').click()